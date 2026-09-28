"""Institutional-grade Excel workbook export engine using openpyxl.

Generates a structured, multi-sheet financial model workbook covering executive
summaries, historical statements, forecast schedules, WACC, DCF valuation,
scenario comparisons, sensitivity matrices, and audit disclosures.
"""

from __future__ import annotations

import io
import logging
from typing import Any, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.reporting.models import ReportBundle

logger = logging.getLogger(__name__)

# Style constants
NAVY_HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
LIGHT_BLUE_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
HIGHLIGHT_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
GRAY_HEADER_FILL = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
SECTION_FONT = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
REGULAR_FONT = Font(name="Calibri", size=10, color="1F2937")
ITALIC_MUTED_FONT = Font(name="Calibri", size=9, italic=True, color="6B7280")
INVALID_FONT = Font(name="Calibri", size=9, color="DC2626")

THIN_SIDE = Side(border_style="thin", color="D1D5DB")
THICK_BOTTOM_SIDE = Side(border_style="medium", color="1E3A8A")
DOUBLE_BOTTOM_SIDE = Side(border_style="double", color="111827")

STANDARD_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
HEADER_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THICK_BOTTOM_SIDE)
TOTAL_BORDER = Border(top=THIN_SIDE, bottom=DOUBLE_BOTTOM_SIDE)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def _autofit_columns(ws: openpyxl.worksheet.worksheet.Worksheet, min_width: int = 14, max_width: int = 40) -> None:
    """Adjust worksheet column widths according to content length."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and not val_str.startswith("="):
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, max_width), min_width)


class ExcelReportGenerator:
    """Orchestrates generation of comprehensive multi-tab Excel workbooks."""

    @classmethod
    def generate_workbook(cls, bundle: ReportBundle) -> bytes:
        """Build and return an Excel workbook as bytes."""
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        cfg = bundle.config
        sec = cfg.sections

        # Sheet 1: Executive Summary (always created if overview or summary selected)
        if sec.include_overview or sec.include_executive_summary:
            cls._build_executive_summary_sheet(wb, bundle)

        # Sheet 2: Historical Statements
        if sec.include_historical_analysis and bundle.historical_bundle and bundle.historical_bundle.periods:
            cls._build_historical_sheet(wb, bundle)

        # Sheet 3: Forecast Projections
        if sec.include_forecast_projections and bundle.forecast_result:
            cls._build_forecast_sheet(wb, bundle)

        # Sheet 4: WACC Analysis
        if sec.include_wacc_analysis and bundle.wacc_result:
            cls._build_wacc_sheet(wb, bundle)

        # Sheet 5: DCF Valuation
        if sec.include_dcf_valuation and bundle.dcf_result:
            cls._build_dcf_sheet(wb, bundle)

        # Sheet 6: Scenario Analysis
        if sec.include_scenario_analysis and bundle.scenario_result:
            cls._build_scenarios_sheet(wb, bundle)

        # Sheet 7: Sensitivity & Simulation
        if sec.include_sensitivity_simulation and (bundle.sensitivity_matrix_result or bundle.monte_carlo_result):
            cls._build_sensitivity_sheet(wb, bundle)

        # Sheet 8: Disclosures & Provenance
        if sec.include_disclosures_limitations or sec.include_appendix:
            cls._build_disclosures_sheet(wb, bundle)

        # If no sheets were generated, add a fallback sheet
        if not wb.sheetnames:
            ws = wb.create_sheet(title="Report Overview")
            ws["A1"] = "No report sections were selected for export."

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    @classmethod
    def _build_executive_summary_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create executive summary and report overview worksheet."""
        ws = wb.create_sheet(title="Executive Summary")
        ws.views.sheetView[0].showGridLines = True

        meta = bundle.config.metadata

        # Title Block
        ws["A1"] = meta.title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = meta.subtitle or "AI-Powered DCF Valuation and Sensitivity Engine"
        ws["A2"].font = ITALIC_MUTED_FONT

        ws["A4"] = "Company:"
        ws["B4"] = bundle.company_name + (f" ({bundle.ticker})" if bundle.ticker else "")
        ws["A5"] = "Valuation Project:"
        ws["B5"] = meta.project_name or f"Project #{bundle.project_id}"
        ws["A6"] = "Reporting Currency:"
        ws["B6"] = meta.reporting_currency
        ws["A7"] = "Report Generation Date:"
        ws["B7"] = meta.report_date
        ws["A8"] = "Prepared By:"
        ws["B8"] = meta.prepared_by or "Financial Modeling Analyst"

        for r in range(4, 9):
            ws[f"A{r}"].font = BOLD_FONT
            ws[f"B{r}"].font = REGULAR_FONT

        # Key Valuation Metrics KPI Box
        row = 10
        ws[f"A{row}"] = "VALUATION HIGHLIGHTS & INTRINSIC RESULTS"
        ws[f"A{row}"].font = SECTION_FONT
        row += 1

        headers = ["Valuation Metric", "Value", "Units / Method", "Notes"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = ALIGN_LEFT if col_idx != 2 else ALIGN_RIGHT
            cell.border = HEADER_BORDER
        row += 1

        dcf = bundle.dcf_result
        if dcf and dcf.is_complete:
            metrics = [
                ("Enterprise Value (EV)", dcf.enterprise_value, f"{meta.reporting_currency} Millions", "PV of Forecast UFCF + PV of Terminal Value"),
                ("Less: Net Debt / Bridge Claims", (dcf.enterprise_value - dcf.equity_value) if dcf.equity_value is not None else None, f"{meta.reporting_currency} Millions", "Cash, Debt, Minority Interest, Preferred"),
                ("Equity Value", dcf.equity_value, f"{meta.reporting_currency} Millions", "Enterprise Value adjusted for net non-operating claims"),
                ("Diluted Shares Outstanding", dcf.diluted_shares, "Millions of shares", "Common share count base"),
                ("Implied Value per Share", dcf.implied_share_price, f"{meta.reporting_currency} per share", "Equity Value / Diluted Common Shares"),
                ("Discount Rate (WACC)", dcf.wacc * 100.0 if dcf.wacc else None, "Percentage (%)", "Weighted Average Cost of Capital"),
                ("Terminal Method", dcf.terminal_inputs.method.replace("_", " ").title(), "Methodology", f"Perpetual g: {dcf.terminal_inputs.perpetual_growth_rate}% | Exit Multiple: {dcf.terminal_inputs.exit_multiple}x"),
            ]
            for label, val, units, notes in metrics:
                ws.cell(row=row, column=1, value=label).font = BOLD_FONT
                val_cell = ws.cell(row=row, column=2, value=val if isinstance(val, (int, float)) else str(val or "N/A"))
                if isinstance(val, float):
                    if "per share" in units:
                        val_cell.number_format = "$#,##0.00"
                    elif "Percentage" in units:
                        val_cell.number_format = "0.00\"%\""
                    else:
                        val_cell.number_format = "$#,##0.00"
                val_cell.font = BOLD_FONT if "Implied" in label or "Equity" in label else REGULAR_FONT
                val_cell.alignment = ALIGN_RIGHT
                ws.cell(row=row, column=3, value=units).font = REGULAR_FONT
                ws.cell(row=row, column=4, value=notes).font = ITALIC_MUTED_FONT

                for c in range(1, 5):
                    ws.cell(row=row, column=c).border = STANDARD_BORDER
                row += 1
        else:
            ws.cell(row=row, column=1, value="DCF valuation is incomplete or unselected.").font = ITALIC_MUTED_FONT
            row += 1

        # Narrative Notes
        if meta.narrative_summary or meta.valuation_notes:
            row += 2
            ws[f"A{row}"] = "EXECUTIVE & METHODOLOGY NOTES"
            ws[f"A{row}"].font = SECTION_FONT
            row += 1
            if meta.narrative_summary:
                ws[f"A{row}"] = "Executive Summary Notes:"
                ws[f"A{row}"].font = BOLD_FONT
                ws[f"B{row}"] = meta.narrative_summary
                ws[f"B{row}"].font = REGULAR_FONT
                row += 1
            if meta.valuation_notes:
                ws[f"A{row}"] = "Valuation & Risk Commentary:"
                ws[f"A{row}"].font = BOLD_FONT
                ws[f"B{row}"] = meta.valuation_notes
                ws[f"B{row}"].font = REGULAR_FONT
                row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_historical_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create historical financial statements worksheet."""
        ws = wb.create_sheet(title="Historical Statements")
        ws.views.sheetView[0].showGridLines = True
        hbundle = bundle.historical_bundle
        if not hbundle:
            return

        ws["A1"] = f"{bundle.company_name} — Historical Financial Analysis"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"Reporting Currency: {bundle.config.metadata.reporting_currency} | Frequency: {hbundle.period_type.title()} | Classification: {hbundle.data_classification}"
        ws["A2"].font = ITALIC_MUTED_FONT

        row = 4
        # Columns: Line Item, Period 1, Period 2, ...
        periods = [p.label for p in hbundle.periods]
        headers = ["Financial Metric / Line Item"] + periods

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = ALIGN_LEFT if col_idx == 1 else ALIGN_RIGHT
            cell.border = HEADER_BORDER
        row += 1

        items = [
            ("Revenue", "revenue", "$#,##0.0"),
            ("Revenue Growth (%)", "revenue_growth_yoy", "0.0%"),
            ("Gross Profit", "gross_profit", "$#,##0.0"),
            ("Gross Margin (%)", "gross_margin", "0.0%"),
            ("EBITDA", "ebitda", "$#,##0.0"),
            ("EBITDA Margin (%)", "ebitda_margin", "0.0%"),
            ("Operating Income (EBIT)", "operating_income", "$#,##0.0"),
            ("EBIT Margin (%)", "operating_margin", "0.0%"),
            ("Net Income", "net_income", "$#,##0.0"),
            ("Operating Cash Flow (CFO)", "operating_cash_flow", "$#,##0.0"),
            ("Capital Expenditures (CapEx)", "capital_expenditures", "$#,##0.0"),
            ("CapEx Intensity (% Revenue)", "capex_pct_revenue", "0.0%"),
            ("CFO Less CapEx", "cfo_less_capex", "$#,##0.0"),
            ("Free Cash Flow (UFCF Estimate)", "ufcf_estimate", "$#,##0.0"),
            ("Operating Net Working Capital (NWC)", "operating_nwc", "$#,##0.0"),
            ("Days Sales Outstanding (DSO)", "dso", "0.0"),
            ("Days Inventory Outstanding (DIO)", "dio", "0.0"),
            ("Days Payables Outstanding (DPO)", "dpo", "0.0"),
            ("Cash Conversion Cycle (CCC)", "cash_conversion_cycle", "0.0"),
        ]

        for label, metric_code, num_fmt in items:
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if "Revenue" in label or "EBITDA" in label or "Free Cash Flow" in label else REGULAR_FONT
            ws.cell(row=row, column=1).border = STANDARD_BORDER

            for p_idx, p_label in enumerate(periods, start=2):
                cell = ws.cell(row=row, column=p_idx)
                cell.border = STANDARD_BORDER
                cell.alignment = ALIGN_RIGHT

                # Lookup value
                val = None
                p_metrics = hbundle.metrics_by_period.get(p_label, {})
                if metric_code in p_metrics:
                    val = p_metrics[metric_code].value
                elif metric_code in ("dso", "dio", "dpo", "cash_conversion_cycle", "operating_nwc"):
                    wc = hbundle.working_capital.get(p_label)
                    if wc:
                        val = getattr(wc, metric_code, None)
                elif metric_code in ("operating_cash_flow", "capital_expenditures", "cfo_less_capex", "ufcf_estimate", "capex_pct_revenue"):
                    cf = hbundle.cash_flow_metrics.get(p_label)
                    if cf:
                        if metric_code == "capital_expenditures":
                            val = cf.capex_magnitude
                        else:
                            val = getattr(cf, metric_code, None)

                if val is not None:
                    # In percentage formatting in bundle, rates are often percentages 0..100, normalize if needed
                    if num_fmt == "0.0%" and abs(val) > 1.0:
                        val = val / 100.0
                    cell.value = float(val)
                    cell.number_format = num_fmt
                else:
                    cell.value = "-"
                    cell.font = ITALIC_MUTED_FONT

            row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_forecast_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create forecast projections and UFCF derivation worksheet."""
        ws = wb.create_sheet(title="Forecast Projections")
        ws.views.sheetView[0].showGridLines = True
        fc = bundle.forecast_result
        if not fc:
            return

        ws["A1"] = f"{bundle.company_name} — Financial Forecast Projections"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"Horizon: {len(fc.projected_years)} Years | Base Period: {fc.base_period_label}"
        ws["A2"].font = ITALIC_MUTED_FONT

        row = 4
        headers = ["Line Item / Derivation Step"] + fc.projected_years
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.alignment = ALIGN_LEFT if col_idx == 1 else ALIGN_RIGHT
            cell.border = HEADER_BORDER
        row += 1

        lines = [
            ("Revenue", fc.revenue, "$#,##0.0"),
            ("Revenue Growth (%)", [r / 100.0 if r else 0.0 for r in fc.revenue_growth_rate], "0.0%"),
            ("Gross Profit", fc.gross_profit, "$#,##0.0"),
            ("Gross Profit Margin (%)", [m / 100.0 if m else 0.0 for m in fc.gross_profit_margin], "0.0%"),
            ("EBITDA", fc.ebitda, "$#,##0.0"),
            ("EBITDA Margin (%)", [m / 100.0 if m else 0.0 for m in fc.ebitda_margin], "0.0%"),
            ("Depreciation & Amortization (D&A)", fc.depreciation_and_amortization, "$#,##0.0"),
            ("Operating Income (EBIT)", fc.ebit, "$#,##0.0"),
            ("EBIT Margin (%)", [m / 100.0 if m else 0.0 for m in fc.ebit_margin], "0.0%"),
            ("Effective Tax Rate (%)", [t / 100.0 if t else 0.0 for t in fc.effective_tax_rate], "0.0%"),
            ("Taxes on Operating Profit", fc.tax_expense, "$#,##0.0"),
            ("NOPAT (EBIT x (1 - t))", fc.nopat, "$#,##0.0"),
            ("(+) Depreciation & Amortization", fc.depreciation_and_amortization, "$#,##0.0"),
            ("(-) Capital Expenditures (CapEx)", [-c for c in fc.capex], "$#,##0.0"),
            ("(-) Change in Operating NWC", [-d for d in fc.delta_operating_nwc], "$#,##0.0"),
            ("(=) Unlevered Free Cash Flow (UFCF)", fc.ufcf, "$#,##0.0"),
        ]

        for label, val_series, num_fmt in lines:
            is_bold = "(=)" in label or "NOPAT" in label or "Revenue" in label or "EBITDA" in label
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if is_bold else REGULAR_FONT
            ws.cell(row=row, column=1).border = TOTAL_BORDER if "(=)" in label else STANDARD_BORDER
            if "(=)" in label:
                ws.cell(row=row, column=1).fill = LIGHT_BLUE_FILL

            for c_idx, val in enumerate(val_series, start=2):
                cell = ws.cell(row=row, column=c_idx, value=float(val))
                cell.number_format = num_fmt
                cell.font = BOLD_FONT if is_bold else REGULAR_FONT
                cell.alignment = ALIGN_RIGHT
                cell.border = TOTAL_BORDER if "(=)" in label else STANDARD_BORDER
                if "(=)" in label:
                    cell.fill = LIGHT_BLUE_FILL

            row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_wacc_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create WACC and capital cost analysis worksheet."""
        ws = wb.create_sheet(title="WACC Analysis")
        ws.views.sheetView[0].showGridLines = True
        wacc = bundle.wacc_result
        if not wacc:
            return

        ws["A1"] = f"{bundle.company_name} — Weighted Average Cost of Capital (WACC)"
        ws["A1"].font = TITLE_FONT

        row = 4
        ws.cell(row=row, column=1, value="WACC COMPONENT & CAPITAL STRUCTURE").font = SECTION_FONT
        row += 1

        headers = ["Parameter / Component", "Value", "Unit / Basis", "Source & Methodology"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.border = HEADER_BORDER
        row += 1

        w_rows = [
            ("Risk-Free Rate (Rf)", wacc.cost_of_equity.risk_free_rate, "0.00%", "Benchmark government bond yield"),
            ("Equity Beta", wacc.cost_of_equity.beta, "0.00x", "Systematic risk exposure proxy"),
            ("Equity Risk Premium (ERP)", wacc.cost_of_equity.equity_risk_premium, "0.00%", "Expected market return over risk-free rate"),
            ("Cost of Equity (Ke = Rf + Beta x ERP)", wacc.cost_of_equity.cost_of_equity, "0.00%", "Capital Asset Pricing Model (CAPM)"),
            ("Pre-tax Cost of Debt (Kd)", wacc.cost_of_debt.pre_tax_cost_of_debt, "0.00%", "Effective borrowing cost or synthetic yield"),
            ("Marginal Tax Rate (t)", wacc.cost_of_debt.effective_tax_rate, "0.00%", "Statutory or forecast tax shield rate"),
            ("After-tax Cost of Debt (Kd x (1 - t))", wacc.cost_of_debt.after_tax_cost_of_debt, "0.00%", "Tax-adjusted debt service cost"),
            ("Market Value of Equity (E)", wacc.capital_structure.equity_value, "$#,##0.0", "Market capitalization or proxy"),
            ("Value of Total Debt (D)", wacc.capital_structure.total_debt, "$#,##0.0", "Short-term + Long-term interest-bearing debt"),
            ("Total Capital (E + D)", wacc.capital_structure.total_capital, "$#,##0.0", "Enterprise financial base"),
            ("Equity Weight (We = E / (E + D))", wacc.capital_structure.weight_equity, "0.0%", "Proportion of equity financing"),
            ("Debt Weight (Wd = D / (E + D))", wacc.capital_structure.weight_debt, "0.0%", "Proportion of debt financing"),
            ("Weighted Average Cost of Capital (WACC)", wacc.wacc, "0.00%", "(We x Ke) + (Wd x Kd_after)"),
        ]

        for label, val, num_fmt, note in w_rows:
            is_wacc = "Weighted Average Cost of Capital" in label
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if is_wacc else REGULAR_FONT
            ws.cell(row=row, column=1).border = TOTAL_BORDER if is_wacc else STANDARD_BORDER
            if is_wacc:
                ws.cell(row=row, column=1).fill = HIGHLIGHT_FILL

            val_cell = ws.cell(row=row, column=2, value=float(val) if val is not None else 0.0)
            val_cell.number_format = num_fmt
            val_cell.font = BOLD_FONT if is_wacc else REGULAR_FONT
            val_cell.alignment = ALIGN_RIGHT
            val_cell.border = TOTAL_BORDER if is_wacc else STANDARD_BORDER
            if is_wacc:
                val_cell.fill = HIGHLIGHT_FILL

            unit_cell = ws.cell(row=row, column=3, value=num_fmt)
            unit_cell.font = REGULAR_FONT
            unit_cell.border = TOTAL_BORDER if is_wacc else STANDARD_BORDER

            note_cell = ws.cell(row=row, column=4, value=note)
            note_cell.font = ITALIC_MUTED_FONT
            note_cell.border = TOTAL_BORDER if is_wacc else STANDARD_BORDER

            row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_dcf_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create DCF valuation, discounting schedule, and equity bridge worksheet."""
        ws = wb.create_sheet(title="DCF Valuation")
        ws.views.sheetView[0].showGridLines = True
        dcf = bundle.dcf_result
        if not dcf:
            return

        ws["A1"] = f"{bundle.company_name} — Discounted Cash Flow Valuation"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"Terminal Method: {dcf.terminal_inputs.method.replace('_', ' ').title()} | Timing: {dcf.discounting_convention.replace('_', ' ').title()}"
        ws["A2"].font = ITALIC_MUTED_FONT

        # 1. Cash Flow Discounting Schedule
        row = 4
        ws.cell(row=row, column=1, value="FORECAST CASH FLOW DISCOUNTING SCHEDULE").font = SECTION_FONT
        row += 1

        headers = ["Period Label", "Projection Year (t)", "Discount Factor", "Forecast UFCF", "Present Value of UFCF"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.border = HEADER_BORDER
        row += 1

        for cf in dcf.cash_flow_schedule:
            ws.cell(row=row, column=1, value=cf.period_label).font = BOLD_FONT
            ws.cell(row=row, column=2, value=cf.period_index).alignment = ALIGN_CENTER
            ws.cell(row=row, column=3, value=cf.discount_factor).number_format = "0.0000"
            ws.cell(row=row, column=4, value=cf.ufcf).number_format = "$#,##0.0"
            ws.cell(row=row, column=5, value=cf.present_value).number_format = "$#,##0.0"

            for c in range(1, 6):
                ws.cell(row=row, column=c).border = STANDARD_BORDER
            row += 1

        # PV of Forecast Cash Flows Total
        ws.cell(row=row, column=1, value="PV of Discrete Forecast UFCF").font = BOLD_FONT
        ws.cell(row=row, column=5, value=dcf.pv_forecast_cash_flows).number_format = "$#,##0.0"
        ws.cell(row=row, column=5).font = BOLD_FONT
        for c in range(1, 6):
            ws.cell(row=row, column=c).border = TOTAL_BORDER
            ws.cell(row=row, column=c).fill = LIGHT_BLUE_FILL
        row += 2

        # 2. Terminal Value Calculation
        ws.cell(row=row, column=1, value="TERMINAL ENTERPRISE VALUE CALCULATION").font = SECTION_FONT
        row += 1

        tv_items = [
            ("Terminal Value Method", dcf.terminal_inputs.method.replace("_", " ").title()),
            ("Terminal Perpetual Growth Rate (g)", f"{dcf.terminal_inputs.perpetual_growth_rate:.2f}%" if dcf.terminal_inputs.perpetual_growth_rate else "N/A"),
            ("Terminal Exit Multiple", f"{dcf.terminal_inputs.exit_multiple:.1f}x EBITDA" if dcf.terminal_inputs.exit_multiple else "N/A"),
            ("Terminal Value (Nominal TV_N)", dcf.terminal_value),
            ("Terminal Value Discount Factor", dcf.terminal_discount_factor),
            ("Present Value of Terminal Value", dcf.pv_terminal_value),
            ("Terminal Value % of Enterprise Value", f"{dcf.terminal_value_pct_of_ev:.1f}%" if dcf.terminal_value_pct_of_ev else "N/A"),
        ]

        for label, val in tv_items:
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if "Present Value" in label else REGULAR_FONT
            vcell = ws.cell(row=row, column=2, value=val if isinstance(val, (int, float, str)) else "N/A")
            if isinstance(val, float):
                vcell.number_format = "$#,##0.0" if "Present Value" in label or "Nominal" in label else "0.0000"
            vcell.font = BOLD_FONT if "Present Value" in label else REGULAR_FONT
            vcell.alignment = ALIGN_RIGHT
            row += 1

        row += 1

        # 3. Enterprise-to-Equity Bridge
        ws.cell(row=row, column=1, value="ENTERPRISE VALUE TO EQUITY VALUE BRIDGE").font = SECTION_FONT
        row += 1

        b = dcf.bridge_inputs
        bridge_rows = [
            ("Enterprise Value (PV Forecast UFCF + PV TV)", dcf.enterprise_value, "+", "$#,##0.0"),
            ("(+) Cash and Cash Equivalents", b.cash_and_equivalents, "+", "$#,##0.0"),
            ("(-) Interest-Bearing Debt", -b.debt_value, "-", "$#,##0.0"),
            ("(-) Minority Interest (Non-Controlling)", -b.minority_interest, "-", "$#,##0.0"),
            ("(-) Preferred Stock Equity", -b.preferred_equity, "-", "$#,##0.0"),
            ("(+/-) Other Net Adjustments", b.other_adjustments, "+/-", "$#,##0.0"),
            ("(=) Implied Equity Value", dcf.equity_value, "=", "$#,##0.0"),
            ("(/) Diluted Common Shares Outstanding", dcf.diluted_shares, "/", "#,##0.0"),
            ("(=) Implied Intrinsic Value per Share", dcf.implied_share_price, "=", "$#,##0.00"),
        ]

        for label, val, sign, num_fmt in bridge_rows:
            is_final = "Implied Intrinsic Value" in label or "Implied Equity Value" in label
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if is_final else REGULAR_FONT
            ws.cell(row=row, column=1).border = TOTAL_BORDER if is_final else STANDARD_BORDER
            if is_final:
                ws.cell(row=row, column=1).fill = HIGHLIGHT_FILL

            vcell = ws.cell(row=row, column=2, value=float(val) if val is not None else "N/A")
            if isinstance(val, float):
                vcell.number_format = num_fmt
            vcell.font = BOLD_FONT if is_final else REGULAR_FONT
            vcell.alignment = ALIGN_RIGHT
            vcell.border = TOTAL_BORDER if is_final else STANDARD_BORDER
            if is_final:
                vcell.fill = HIGHLIGHT_FILL

            row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_scenarios_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create Base/Bull/Bear scenario comparison worksheet."""
        ws = wb.create_sheet(title="Scenario Analysis")
        ws.views.sheetView[0].showGridLines = True
        sc = bundle.scenario_result
        if not sc:
            return

        ws["A1"] = f"{bundle.company_name} — Scenario Analysis (Base / Bull / Bear)"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"Scenario Set: {bundle.scenario_model_name or 'Default Scenarios'}"
        ws["A2"].font = ITALIC_MUTED_FONT

        row = 4
        headers = ["Valuation & Operating Driver", "Base Case", "Bull Case", "Bear Case"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=h)
            cell.font = WHITE_BOLD_FONT
            cell.fill = NAVY_HEADER_FILL
            cell.border = HEADER_BORDER
        row += 1

        b_case = sc.base_case
        u_case = sc.bull_case
        d_case = sc.bear_case

        s_rows = [
            ("Revenue Growth Override (pp)", b_case.revenue_growth_override_pp, u_case.revenue_growth_override_pp, d_case.revenue_growth_override_pp, "0.0\" pp\""),
            ("Operating Margin Override (pp)", b_case.margin_override_pp, u_case.margin_override_pp, d_case.margin_override_pp, "0.0\" pp\""),
            ("WACC Adjustment (bps)", b_case.wacc_adjustment_bps, u_case.wacc_adjustment_bps, d_case.wacc_adjustment_bps, "0\" bps\""),
            ("Effective Discount Rate (WACC)", b_case.effective_wacc, u_case.effective_wacc, d_case.effective_wacc, "0.00%"),
            ("Enterprise Value (EV)", b_case.enterprise_value, u_case.enterprise_value, d_case.enterprise_value, "$#,##0.0"),
            ("Equity Value", b_case.equity_value, u_case.equity_value, d_case.equity_value, "$#,##0.0"),
            ("Implied Share Price", b_case.implied_share_price, u_case.implied_share_price, d_case.implied_share_price, "$#,##0.00"),
            ("Upside / Downside vs. Base (%)", 0.0, u_case.upside_downside_pct, d_case.upside_downside_pct, "0.0%"),
            ("Validation Status", b_case.status, u_case.status, d_case.status, "@"),
        ]

        for label, b_val, u_val, d_val, num_fmt in s_rows:
            is_bold = "Implied Share Price" in label or "Enterprise Value" in label
            ws.cell(row=row, column=1, value=label).font = BOLD_FONT if is_bold else REGULAR_FONT
            ws.cell(row=row, column=1).border = STANDARD_BORDER

            for c_idx, val in enumerate([b_val, u_val, d_val], start=2):
                cell = ws.cell(row=row, column=c_idx)
                cell.border = STANDARD_BORDER
                cell.alignment = ALIGN_RIGHT
                if isinstance(val, float):
                    if num_fmt == "0.0%" and abs(val) > 1.0:
                        val = val / 100.0
                    cell.value = val
                    cell.number_format = num_fmt
                    cell.font = BOLD_FONT if is_bold else REGULAR_FONT
                else:
                    cell.value = str(val or "N/A")
                    cell.font = REGULAR_FONT

            row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_sensitivity_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create 2D sensitivity matrix and Monte Carlo simulation worksheet."""
        ws = wb.create_sheet(title="Sensitivity & Simulation")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = f"{bundle.company_name} — Valuation Sensitivity & Simulation"
        ws["A1"].font = TITLE_FONT

        row = 4
        # 1. 2D Sensitivity Matrix
        matrix = bundle.sensitivity_matrix_result
        if matrix:
            ws.cell(row=row, column=1, value=f"TWO-DIMENSIONAL SENSITIVITY MATRIX ({matrix.metric_label})").font = SECTION_FONT
            row += 1

            col_headers = [f"WACC \\ {matrix.col_label}"] + [f"{c:.2f}%" if "%" in matrix.col_label else f"{c:.1f}x" for c in matrix.col_values]
            for col_idx, ch in enumerate(col_headers, start=1):
                cell = ws.cell(row=row, column=col_idx, value=ch)
                cell.font = WHITE_BOLD_FONT
                cell.fill = NAVY_HEADER_FILL
                cell.border = HEADER_BORDER
                cell.alignment = ALIGN_CENTER
            row += 1

            for r_idx, r_val in enumerate(matrix.row_values):
                ws.cell(row=row, column=1, value=f"{r_val:.2f}%").font = BOLD_FONT
                ws.cell(row=row, column=1).border = STANDARD_BORDER
                ws.cell(row=row, column=1).alignment = ALIGN_CENTER

                for c_idx, c_val in enumerate(matrix.col_values):
                    cell_res = matrix.cells[r_idx][c_idx]
                    cell = ws.cell(row=row, column=c_idx + 2)
                    cell.border = STANDARD_BORDER
                    cell.alignment = ALIGN_RIGHT

                    if cell_res.is_valid and cell_res.output_value is not None:
                        cell.value = cell_res.output_value
                        cell.number_format = "$#,##0.00" if "Share" in matrix.metric_label else "$#,##0.0"
                        if cell_res.is_baseline:
                            cell.fill = HIGHLIGHT_FILL
                            cell.font = BOLD_FONT
                    else:
                        cell.value = "INVALID"
                        cell.font = INVALID_FONT
                        cell.fill = GRAY_HEADER_FILL

                row += 1

            row += 2

        # 2. Monte Carlo Simulation Summary
        mc = bundle.monte_carlo_result
        if mc and mc.summary_stats:
            ws.cell(row=row, column=1, value="MONTE CARLO PROBABILISTIC SIMULATION SUMMARY").font = SECTION_FONT
            row += 1

            ws.cell(row=row, column=1, value=f"Total Iterations: {mc.total_iterations} | Valid Draws: {mc.valid_iterations} | Invalid Draws: {mc.invalid_iterations} | Random Seed: {mc.config.random_seed}").font = ITALIC_MUTED_FONT
            row += 1

            headers = ["Statistic / Percentile", "Enterprise Value", "Equity Value", "Implied Share Price"]
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col_idx, value=h)
                cell.font = WHITE_BOLD_FONT
                cell.fill = SUBHEADER_FILL
                cell.border = HEADER_BORDER
            row += 1

            ev_s = mc.summary_stats.get("enterprise_value")
            eq_s = mc.summary_stats.get("equity_value")
            sp_s = mc.summary_stats.get("implied_share_price")

            stat_keys = [
                ("Mean", "mean"),
                ("Median (P50)", "median"),
                ("Std Deviation", "std_dev"),
                ("Minimum", "min_val"),
                ("10th Percentile (P10)", "p10"),
                ("25th Percentile (P25)", "p25"),
                ("75th Percentile (P75)", "p75"),
                ("90th Percentile (P90)", "p90"),
                ("Maximum", "max_val"),
            ]

            for label, attr in stat_keys:
                ws.cell(row=row, column=1, value=label).font = BOLD_FONT if "Median" in label or "Mean" in label else REGULAR_FONT
                ws.cell(row=row, column=1).border = STANDARD_BORDER

                for c_idx, s in enumerate([ev_s, eq_s, sp_s], start=2):
                    cell = ws.cell(row=row, column=c_idx)
                    cell.border = STANDARD_BORDER
                    cell.alignment = ALIGN_RIGHT
                    val = getattr(s, attr, None) if s else None
                    if val is not None:
                        cell.value = val
                        cell.number_format = "$#,##0.00" if c_idx == 4 else "$#,##0.0"
                    else:
                        cell.value = "-"
                row += 1

        _autofit_columns(ws)

    @classmethod
    def _build_disclosures_sheet(cls, wb: openpyxl.Workbook, bundle: ReportBundle) -> None:
        """Create data quality disclosures and modeling limitations worksheet."""
        ws = wb.create_sheet(title="Audit & Disclosures")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = f"{bundle.company_name} — Disclosures, Data Quality & Model Audit"
        ws["A1"].font = TITLE_FONT

        row = 4
        # Missing Models & Validation Diagnostics
        ws.cell(row=row, column=1, value="MODEL VALIDATION & DATA INTEGRITY FINDINGS").font = SECTION_FONT
        row += 1

        if bundle.missing_models_diagnostics:
            for diag in bundle.missing_models_diagnostics:
                ws.cell(row=row, column=1, value="⚠️ Missing Dependency:").font = BOLD_FONT
                ws.cell(row=row, column=2, value=diag).font = REGULAR_FONT
                row += 1

        if bundle.validation_warnings:
            for warn in bundle.validation_warnings:
                ws.cell(row=row, column=1, value="⚠️ Calculation Warning:").font = BOLD_FONT
                ws.cell(row=row, column=2, value=warn).font = REGULAR_FONT
                row += 1

        if bundle.data_quality_issues:
            for dq in bundle.data_quality_issues:
                ws.cell(row=row, column=1, value="ℹ️ Data Quality Note:").font = BOLD_FONT
                ws.cell(row=row, column=2, value=dq).font = REGULAR_FONT
                row += 1

        if not bundle.missing_models_diagnostics and not bundle.validation_warnings and not bundle.data_quality_issues:
            ws.cell(row=row, column=1, value="No blocking validation errors or material data quality issues detected.").font = REGULAR_FONT
            row += 1

        row += 2
        # Analytical Disclaimers
        ws.cell(row=row, column=1, value="ANALYTICAL DISCLAIMERS & IMPORTANT NOTICES").font = SECTION_FONT
        row += 1

        disclaimers = [
            "1. Not Investment Advice: This report is generated solely for analytical, educational, and illustrative financial modeling purposes. It does not constitute investment advice, a financial promotion, or an offer or recommendation to buy or sell securities.",
            "2. User-Selected Inputs: All valuation conclusions, forecast trajectories, and scenario projections are strictly conditional on user-specified assumptions and historical statements entered into the platform.",
            "3. Simulation Limitations: Monte Carlo distributions and sensitivity tables reflect statistical variations under specified parameters and do not represent probability distributions of real-world equity market returns.",
            "4. Model Risk: Intrinsic value calculations depend heavily on terminal value assumptions, cost of capital estimates, and multi-year forecasting accuracy. Actual operating performance and market conditions may differ materially.",
        ]

        for disc in disclaimers:
            ws.cell(row=row, column=1, value=disc).font = ITALIC_MUTED_FONT
            row += 1

        _autofit_columns(ws, max_width=80)
